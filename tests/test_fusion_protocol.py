from __future__ import annotations

import copy
import csv
import json
import random
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest
import torch
from openpyxl import load_workbook

from models.identity_fusion_data import IdentitySequenceHomologyShardDataset, load_identity_sidecar
from models.sequence_homology_fusion_baselines import (
    build_sequence_homology_fixed_fusion_model, build_sequence_homology_identity_fusion_model,
    identity_neural_weight,
)
from reliability_aware.utils import model_training as training
from reliability_aware.utils.fusion_preparation import prepare_identity_split
from reliability_aware.utils.fusion_protocol import (
    BASELINES, VALIDATION_EXCLUSIONS, bind_run_directory, selected_candidates,
    validation_keep_ids, fusion_parameters,
)
from reliability_aware.utils.inference_runtime import load_model_from_checkpoint
from reliability_aware.utils.prediction_cache import canonical_ids_hash

HP = dict(learning_rate=.001, weight_decay=.0001, attn_hidden_dim=4, attn_dropout=0,
          head_hidden_dim=4, head_dropout=0, identity_a=.25, identity_k=5.,
          pos_weight_cap=20., lambda_hier=0.)


def toy_resources(tmp_path):
    fasta = tmp_path / 'val.fasta'
    fasta.write_text(''.join(f'>{q}\n' + 'A' * 100 + '\n' for q in ['good', '2VAU-A', 'nohit', '5LSQ-A']))
    valid = 'good\ts1\t1e-9\t100\t50\t100\t100\t50\t50'
    other = 'good\ts2\t1e-9\t50\t50\t100\t100\t80\t80'
    original, enriched = tmp_path / 'old.tsv', tmp_path / 'new.tsv'
    original.write_text(valid + '\n' + other + '\n2VAU-A\tbad\t1e-9\t100\t80\t100\t66\t99\t100\n')
    enriched.write_text(valid + '\t50\n' + other + '\t80\n2VAU-A\tbad\t1e-9\t100\t80\t100\t66\t99\t100\t99\n5LSQ-A\tbad\t1e-9\t100\t80\t100\t66\t71\t100\t71\n')
    return {'prepared_dir': str(tmp_path / 'prepared'), 'validation_exclude_ids': sorted(VALIDATION_EXCLUSIONS),
            'pdb_val': {'fasta': str(fasta), 'original_hits': str(original), 'enriched_hits': str(enriched),
                        'identity': str(tmp_path / 'prepared/identity.json')}}


def test_exclusion_removes_entire_queries_preserves_sources_and_reuses(tmp_path):
    resources = toy_resources(tmp_path)
    source = Path(resources['pdb_val']['enriched_hits'])
    before = source.read_bytes()
    report = prepare_identity_split(resources=resources, split='pdb_val')
    assert report['status'] == 'matched'
    assert report['remaining_query_count'] == 2
    assert len(report['removed_enriched_rows']) == 2
    sidecar = load_identity_sidecar(resources['pdb_val']['identity'])
    assert set(sidecar) == {'good', 'nohit'}
    assert sidecar['good']['identity_fraction'] == pytest.approx(.6)
    assert sidecar['nohit']['identity_fraction'] is None
    assert source.read_bytes() == before
    assert prepare_identity_split(resources=resources, split='pdb_val')['status'] == 'matched'
    for model in BASELINES:
        assert validation_keep_ids(model, {'good', *VALIDATION_EXCLUSIONS}, VALIDATION_EXCLUSIONS) == {'good'}
    legacy = {'good', *VALIDATION_EXCLUSIONS}
    assert validation_keep_ids('sequence_homology_confidence_gate', legacy, VALIDATION_EXCLUSIONS) is legacy


@pytest.mark.parametrize('mutation', ['bitscore', 'conflicting_count', 'third_corrupt'])
def test_enrichment_rejects_mismatch_and_ambiguity(tmp_path, mutation):
    resources = toy_resources(tmp_path)
    p = Path(resources['pdb_val']['enriched_hits'])
    if mutation == 'bitscore':
        p.write_text(p.read_text().replace('good\ts1\t1e-9\t100', 'good\ts1\t1e-9\t101'))
    elif mutation == 'conflicting_count':
        p.write_text(p.read_text() + 'good\ts1\t1e-9\t100\t50\t100\t100\t50\t50\t49\n')
    else:
        p.write_text(p.read_text() + 'good\tx\t1e-9\t10\t50\t100\t66\t100\t100\t101\n')
    with pytest.raises(ValueError):
        prepare_identity_split(resources=resources, split='pdb_val')
    assert not Path(resources['pdb_val']['identity']).exists()
    report = json.loads(Path(resources['prepared_dir'], 'pdb_val_evidence_report.json').read_text())
    assert report['status'] == 'failed'


def test_loader_filters_indices_and_aligns_shuffled_records(tmp_path):
    resources = toy_resources(tmp_path)
    prepare_identity_split(resources=resources, split='pdb_val')
    identity = Path(resources['pdb_val']['identity'])
    raw = json.loads(identity.read_text()); raw['records'].reverse(); identity.write_text(json.dumps(raw))
    esm, hom = tmp_path / 'esm', tmp_path / 'hom'
    esm.mkdir(); hom.mkdir()
    labels = ['good', '2VAU-A', 'nohit', '5LSQ-A']
    manifest = tmp_path / 'manifest.csv'
    with manifest.open('w', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(['shard_number', 'local_seq_idx', 'global_seq_idx', 'label', 'sequence_length', 'truncated_length'])
        writer.writerows([0, i, i, label, 2, 2] for i, label in enumerate(labels))
    torch.save({'representations': [torch.full((2, 1280), float(i)) for i in range(4)],
                'labels': labels, 'seq_lengths': [2]*4, 'trunc_lengths': [2]*4}, esm / 'part_0000.pt')
    torch.save({'priors': [torch.tensor([float(i)]) for i in range(4)],
                'gate_features': [torch.tensor([0., 0., float(np.log(3)) if i == 0 else 0., float(i != 2)]) for i in range(4)],
                'stats': [{} for _ in labels]}, hom / 'homology_shard_0000.pt')
    ds = IdentitySequenceHomologyShardDataset(esm_shard_dir=esm, homology_shard_dir=hom,
            manifest_path=manifest, keep_ids={'good', 'nohit'}, identity_sidecar_path=identity)
    assert len(ds) == 2
    assert [ds[i]['global_idx'] for i in range(2)] == [0, 2]
    assert ds[0]['identity_fraction'] == pytest.approx(.6)
    assert ds[1]['label'] == 'nohit' and ds[1]['homology_scores'].item() == 2.
    assert torch.equal(ds[1]['rep'], torch.full((2, 1280), 2.))
    raw['records'].append(raw['records'][0]); identity.write_text(json.dumps(raw))
    with pytest.raises(ValueError, match='duplicate'):
        load_identity_sidecar(identity)


@pytest.mark.parametrize('values,epoch,stopped', [
    ([.50, .49, .51, .50, .505, .51], 3, 6),
    ([.50, .49, .51, .50, .52, .51, .51, .51], 5, 8),
])
def test_patience_selects_strict_best_and_resets(tmp_path, monkeypatch, values, epoch, stopped):
    model = torch.nn.Linear(1, 1)
    opt = torch.optim.AdamW(model.parameters())
    metrics = iter(values)
    monkeypatch.setattr(training, 'train_one_epoch', lambda **k: .1)
    monkeypatch.setattr(training, 'evaluate_model', lambda **k: dict(Fmax=next(metrics), AUPR=.2, Smin=.3, val_loss=.1))
    result = training.fit_model(model, [], [], opt, None, None, [], 'BP', None, [], 'cpu',
                                num_epochs=len(values), patience=3, out_dir=tmp_path, hparams={})
    assert result['best_epoch'] == epoch and len(result['val_Fmax']) == stopped
    assert result['consecutive_non_improvement'] == 3
    assert torch.load(tmp_path / 'best_model.pt', weights_only=False)['epoch'] == epoch


def test_nonfinite_validation_cannot_save_checkpoint(tmp_path, monkeypatch):
    model = torch.nn.Linear(1, 1)
    monkeypatch.setattr(training, 'train_one_epoch', lambda **k: .1)
    monkeypatch.setattr(training, 'evaluate_model', lambda **k: dict(Fmax=float('nan'), AUPR=.2, Smin=.3, val_loss=.1))
    with pytest.raises(RuntimeError, match='Non-finite'):
        training.fit_model(model, [], [], torch.optim.AdamW(model.parameters()), None, None, [], 'BP', None, [], 'cpu', out_dir=tmp_path)
    assert not (tmp_path / 'best_model.pt').exists()


@pytest.mark.parametrize('builder', [build_sequence_homology_fixed_fusion_model, build_sequence_homology_identity_fusion_model])
def test_branch_gradients_optimizer_and_checkpoint_roundtrip(tmp_path, builder):
    torch.manual_seed(3)
    model, optimizer = builder(HP, ['g1', 'g2'], 'cpu')
    x, h = torch.randn(2, 3, 1280), torch.rand(2, 2)
    before_x, before_h = x.clone(), h.clone()
    args = dict(padded=x, mask=torch.ones(2, 3, dtype=torch.bool), homology_scores=h,
                identity_fraction=torch.tensor([.2, .4]), has_retained_hit=torch.tensor([True, True]))
    before = copy.deepcopy(model.state_dict())
    out = model(**args, gate_features=torch.zeros(2, 4))
    assert torch.equal(out['probs'], model(**args, gate_features=torch.ones(2, 4))['probs'])
    out['probs'].square().sum().backward()
    for module in [model.seq_branch, model.head]:
        assert any(p.grad is not None and torch.isfinite(p.grad).all() and p.grad.abs().sum() > 0 for p in module.parameters())
    params = [id(p) for group in optimizer.param_groups for p in group['params']]
    assert len(params) == len(set(params)) == len(list(model.parameters()))
    optimizer.step()
    assert torch.equal(x, before_x) and torch.equal(h, before_h)
    if hasattr(model, 'identity_a'):
        assert torch.equal(before['identity_a'], model.identity_a)
        assert torch.equal(before['identity_k'], model.identity_k)
    checkpoint = {'model_state_dict': model.state_dict(), 'hparams': HP,
                  'fusion_protocol_version': 'validation_exclude_two_v1', 'go_terms': ['g1', 'g2'],
                  'go_terms_sha256': canonical_ids_hash(['g1', 'g2']), 'fusion_parameters': fusion_parameters(model)}
    path = tmp_path / 'model.pt'; torch.save(checkpoint, path)
    clone, _ = load_model_from_checkpoint(model_builder=builder, checkpoint_path=path, go_terms=['g1', 'g2'], device='cpu')
    assert torch.equal(model(**args)['probs'], clone(**args)['probs'])
    key = next(iter(checkpoint['fusion_parameters'])); checkpoint['fusion_parameters'][key] += .2
    torch.save(checkpoint, path)
    with pytest.raises(ValueError, match='Tampered'):
        load_model_from_checkpoint(model_builder=builder, checkpoint_path=path, go_terms=['g1', 'g2'], device='cpu')


def test_protocol_rejects_old_cohort_and_sorts_candidates(tmp_path):
    metadata = {'validation_exclude_ids': sorted(VALIDATION_EXCLUSIONS)}
    bind_run_directory(tmp_path / 'search', metadata)
    path = tmp_path / 'search/search_results.json'
    path.write_text(json.dumps([{'score': .1, 'hparams': {'identity_a': .1}}, {'score': .8, 'hparams': {'identity_a': .8}}]))
    assert selected_candidates(path, 1, metadata) == [{'identity_a': .8}]
    with pytest.raises(ValueError, match='different cohort'):
        selected_candidates(path, 1, {'validation_exclude_ids': []})
    with pytest.raises(ValueError, match='Protocol mismatch'):
        bind_run_directory(path.parent, {'validation_exclude_ids': []})


def test_final_handoff_saves_actual_best_weights(tmp_path):
    observed = []
    def builder(hp, terms, device):
        observed.append(hp['identity_a'])
        model = torch.nn.Linear(1, 1)
        return model, torch.optim.AdamW(model.parameters())
    def fit(**kw):
        out = Path(kw['out_dir']); out.mkdir(parents=True, exist_ok=True)
        score = kw['hparams']['identity_a']
        torch.save({'model_state_dict': kw['model'].state_dict(), 'hparams': kw['hparams']}, out / 'best_model.pt')
        return {'val_Fmax': [score], 'val_AUPR': [.1], 'val_Smin': [.2]}
    result = training.run_model_training([dict(HP, identity_a=.8), dict(HP, identity_a=.2)],
            train_loader=[], val_loader=[], train_keep_ids_for_aspect={'train'}, train_label_to_indices={'train': [0]},
            go_terms=['g'], child_parent_pairs=torch.empty((0, 2), dtype=torch.long), go_aspect='BP', obo_path=None,
            train_annotations=[], build_model_fn=builder, fit_function=fit, device='cpu', base_dir=tmp_path,
            seed=42, checkpoint_extra={})
    saved = torch.load(tmp_path / 'best_model.pt', weights_only=False)
    assert observed == [.8, .2] and result['run_id'] == 0
    assert saved['model_state_dict'] and saved['hparams']['identity_a'] == .8
    assert saved['hparams']['training_seed'] == 10042


def test_rule_rejects_nonfinite_parameters():
    for k in [float('nan'), float('inf'), -1]:
        with pytest.raises(ValueError):
            identity_neural_weight(torch.tensor([.5]), .2, k)
    assert torch.equal(identity_neural_weight(torch.tensor([0., .5, 1.]), .2, 0.), torch.ones(3))


def test_search_rng_is_independent_and_reproducible(tmp_path):
    from reliability_aware.utils.model_randomized_search import run_randomized_search
    def builder(hp, terms, device):
        # The trainer resets Python/NumPy/PyTorch streams for every candidate.
        model = torch.nn.Linear(1, 1)
        return model, torch.optim.AdamW(model.parameters())
    def fit(**kw):
        Path(kw['out_dir']).mkdir(parents=True, exist_ok=True)
        random.seed(42)
        return {'val_Fmax': [kw['hparams']['identity_a']], 'val_AUPR': [.1], 'val_Smin': [.2]}
    space = {key: [value] for key, value in HP.items()}
    space['identity_a'] = [0., .2, .4, .6, .8, 1.]
    def run(directory):
        return run_randomized_search(train_keep_ids_for_aspect={'train'}, train_label_to_indices={'train': [0]},
            go_terms=['g'], child_parent_pairs=torch.empty((0, 2), dtype=torch.long), go_aspect='BP', obo_path=None,
            train_annotations=[], search_space=space, device='cpu', num_trials=4, trial_epochs=1,
            train_loader=[], val_loader=[], fit_function=fit, build_model_fn=builder, smoke_test_fn=None,
            seed=42, checkpoint_extra={}, base_dir=directory)
    a, b = run(tmp_path / 'a'), run(tmp_path / 'b')
    assert [r['hparams'] for r in a] == [r['hparams'] for r in b]
    assert len({r['hparams']['identity_a'] for r in a}) == 4
    assert sorted(r['hparams']['training_seed'] for r in a) == [42, 43, 44, 45]


def test_top_five_ties_and_training_self_exclusion(tmp_path):
    from reliability_aware.utils.identity_fusion import build_identity_sidecar
    from reliability_aware.utils.diamond_homology import DiamondSearchConfig
    path = tmp_path / 'hits.tsv'
    rows = ['q\tq\t1e-20\t1000\t100\t100\t100\t100\t100\t100\n']
    rows += [f'q\t{target}\t1e-9\t100\t50\t100\t100\t50\t50\t{count}\n'
             for target, count in [('f', 99), ('e', 50), ('d', 50), ('c', 50), ('b', 50), ('a', 50)]]
    path.write_text(''.join(rows))
    result = build_identity_sidecar(query_ids=['q'], hits_tsv=path, output_path=tmp_path / 'identity.json',
                                   config=DiamondSearchConfig(), exclude_self_hits=True)
    assert result['records'][0]['identity_fraction'] == .5
    assert result['records'][0]['retained_hit_count'] == 6
    assert [x['target_id'] for x in json.loads(Path(result['audit_path']).read_text())] == list('abcde')


def test_uploaded_combiner_parity_without_importing_reference_pipeline():
    import ast
    source = Path(__file__).resolve().parents[2] / 'InterLabelGO-main/InterLabelGO-main/predict.py'
    if not source.exists():
        pytest.skip('Read-only uploaded reference not available in this checkout')
    fn = next(x for x in ast.walk(ast.parse(source.read_text())) if isinstance(x, ast.FunctionDef) and x.name == 'combine_score')
    module = ast.fix_missing_locations(ast.Module(body=[fn], type_ignores=[]))
    env = {'np': np, 'a_weight_dict': {'BPO': .45}, 'k_weight_dict': {'BPO': 1.5}}
    exec(compile(module, str(source), 'exec'), env)
    for s in [0., .2, .6, 1.]:
        expected = env['combine_score'](None, ('q', .7, .2), aspect='BPO', seqid_dict={'q': s}, seqid_combine=True)
        w = identity_neural_weight(torch.tensor([s], dtype=torch.float64), .45, 1.5).item()
        assert .7*w + .2*(1-w) == pytest.approx(expected)


def test_whole_test_paired_bootstrap_and_workbook(tmp_path):
    from reliability_aware.utils.fusion_reporting import compare_pair, write_fusion_workbook
    from reliability_aware.utils.prediction_cache import PredictionCache, sha256_file
    obo = tmp_path / 'go.obo'
    obo.write_text('''format-version: 1.2

[Term]
id: GO:0008150
name: root
namespace: biological_process

[Term]
id: GO:0000001
name: child
namespace: biological_process
is_a: GO:0008150 ! root
''')
    metadata = {'dataset': 'pdb_test', 'go_aspect': 'BP', 'train_annotations': [['GO:0000001']],
                'source_hashes': {'obo': {'sha256': sha256_file(obo)}, 'annotations': {'sha256': 'annotations'},
                                  'test_fasta': {'sha256': 'queries'}}}
    cache = PredictionCache(np.array(['q1', 'q2']), np.array(['GO:0000001']), np.array([[.8], [.5]]),
                            np.ones((2, 1)), np.ones(2, bool), metadata)
    rows, arrays = compare_pair(cache, cache, dataset='pdb_test', aspect='BP', obo=obo,
                               output_dir=tmp_path / 'plans', n_resamples=12)
    assert len(rows) == 3 and all(r['delta'] == 0 and r['ci_lower'] == r['ci_upper'] == 0 for r in rows)
    rows36 = [dict(r, dataset=d, aspect=a, comparator=m) for d in ['pdb_test', 'af_test']
              for a in ['BP', 'MF', 'CC'] for m in ['fixed', 'identity'] for r in rows]
    path = write_fusion_workbook(rows36, {}, {'synthetic_test': True}, tmp_path / 'toy.xlsx')
    book = load_workbook(path, read_only=True)
    assert book['Performance'].max_row == 37
    assert {'PDB', 'AF', 'Fixed_weights', 'Identity_parameters', 'Provenance'} <= set(book.sheetnames)
    book.close()
    bad = copy.deepcopy(cache); bad.metadata['dataset'] = 'af_test'
    with pytest.raises(ValueError, match='dataset mismatch'):
        compare_pair(cache, bad, dataset='pdb_test', aspect='BP', obo=obo, output_dir=tmp_path, n_resamples=2)
