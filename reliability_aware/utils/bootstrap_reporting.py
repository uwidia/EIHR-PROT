"""Retained-hit identity bootstrap orchestration and auditable Excel export."""
from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from reliability_aware.utils.bootstrap import PreparedCafaEvaluator, ResamplingPlan, percentile_interval
from reliability_aware.utils.prediction_cache import PredictionCache, align_caches, read_prediction_cache

CATEGORIES = ["identity_ge95", "identity_70_lt95", "identity_50_lt70", "identity_40_lt50", "identity_30_lt40", "identity_lt30", "no_retained_hit"]
LABELS = {"identity_ge95": "≥95%", "identity_70_lt95": "70–<95%", "identity_50_lt70": "50–<70%", "identity_40_lt50": "40–<50%", "identity_30_lt40": "30–<40%", "identity_lt30": "<30%", "no_retained_hit": "No retained hit"}
METRICS = ["Fmax", "AUPR", "Smin"]


def load_assignments(path: Path) -> dict[str, str]:
    with path.open(newline="") as handle:
        rows = list(csv.DictReader(handle))
    seen: set[str] = set(); values: dict[str, str] = {}
    for row in rows:
        protein, category = row.get("test_protein_id"), row.get("category")
        if not protein or not category or category not in CATEGORIES:
            raise ValueError(f"Invalid retained-hit assignment row: {row}")
        if protein in seen:
            raise ValueError(f"Duplicate assignment ID: {protein}")
        seen.add(protein); values[protein] = category
    return values


def _subcache(cache: PredictionCache, mask: np.ndarray) -> PredictionCache:
    return PredictionCache(cache.protein_ids[mask], cache.go_terms, cache.probabilities[mask], cache.labels[mask], cache.eligibility[mask], cache.metadata, cache.gate_weights[mask] if cache.gate_weights is not None else None)


def _train_annotations(cache: PredictionCache) -> list[set[str]]:
    source = cache.metadata.get("train_annotations")
    if source is None:
        raise ValueError("Cache lacks immutable train_annotations required for checkpoint-free fixed IC")
    return [set(terms) for terms in source]


def _metric_result(prepared: PreparedCafaEvaluator, plan: ResamplingPlan, batch_size: int) -> tuple[dict[str, float], dict[str, np.ndarray]]:
    point = prepared.metrics(np.ones(prepared.precision.shape[1], dtype=np.int32))
    return point, prepared.bootstrap(plan, batch_size=batch_size)


def analyze_pair(*, full_cache_path: Path, ablation_cache_path: Path, aspect: str,
                 assignments: dict[str, str], obo_path: Path, dataset: str,
                 plans_dir: Path, n_resamples: int = 10000, seed: int = 42,
                 batch_size: int = 64) -> dict[str, list[dict[str, Any]]]:
    full, ablation = align_caches(read_prediction_cache(full_cache_path), read_prediction_cache(ablation_cache_path))
    if set(full.protein_ids) - set(assignments):
        raise ValueError("Cache contains proteins absent from assignments.csv")
    raw_counts = Counter(assignments.values())
    records = {"counts": [], "gate": [], "gate_paired": [], "performance": [], "replicates": {}}
    for order, category in enumerate(CATEGORIES, 1):
        mask = np.asarray([(assignments.get(pid) == category) and eligible for pid, eligible in zip(full.protein_ids, full.eligibility)])
        n = int(mask.sum())
        records["counts"].append({"aspect": aspect, "category": category, "display_label": LABELS[category], "raw_n": raw_counts[category], "evaluated_n": n, "excluded_n": raw_counts[category] - n, "exclusion_reason": "not annotation-eligible or not covered by eligible-only cache"})
        if n == 0:
            for model in ("sequence_homology_confidence_gate", "sequence_homology_internal_gate"):
                records["gate"].append({"aspect": aspect, "category": category, "order": order, "display_label": LABELS[category], "n": 0, "model": model, "point": None, "ci_lower": None, "ci_upper": None, "lower_error": None, "upper_error": None, "n_requested": n_resamples, "n_valid": 0, "status": "NA: empty eligible stratum"})
            records["gate_paired"].append({"aspect": aspect, "category": category, "n": 0, "full": None, "ablation": None, "delta": None, "ci_lower": None, "ci_upper": None, "status": "NA: empty eligible stratum"})
            for metric in METRICS:
                records["performance"].append({"aspect": aspect, "category": category, "n": 0, "metric": metric, "full": None, "ablation": None, "delta": None, "ci_lower": None, "ci_upper": None, "status": "NA: empty eligible stratum"})
            continue
        f, a = _subcache(full, mask), _subcache(ablation, mask)
        plan = ResamplingPlan.load_or_create(plans_dir / f"{dataset}_{aspect}_{category}_seed{seed}_B{n_resamples}.npz", dataset=dataset, aspect=aspect, cohort_name=category, protein_ids=f.protein_ids, n_resamples=n_resamples, seed=seed)
        if f.gate_weights is None or a.gate_weights is None:
            raise ValueError("Gate comparison requires gate_weights in both caches")
        fg, ag = f.gate_weights[:, 1], a.gate_weights[:, 1]
        bootstrap_means = lambda values: values[plan.indices].mean(axis=1)
        fgr, agr = bootstrap_means(fg), bootstrap_means(ag)
        for model, values, replicates in (("sequence_homology_confidence_gate", fg, fgr), ("sequence_homology_internal_gate", ag, agr)):
            lo, hi = percentile_interval(replicates)
            point = float(values.mean())
            records["gate"].append({"aspect": aspect, "category": category, "order": order, "display_label": LABELS[category], "n": n, "model": model, "point": point, "ci_lower": lo, "ci_upper": hi, "lower_error": point - lo, "upper_error": hi - point, "n_requested": n_resamples, "n_valid": n_resamples, "status": "ok"})
        delta = fgr - agr; lo, hi = percentile_interval(delta)
        records["gate_paired"].append({"aspect": aspect, "category": category, "n": n, "full": float(fg.mean()), "ablation": float(ag.mean()), "delta": float(fg.mean()-ag.mean()), "ci_lower": lo, "ci_upper": hi, "status": "ok"})
        fp, eligible_f = PreparedCafaEvaluator.prepare(y_true=f.labels, y_prob=f.probabilities, go_terms=f.go_terms.tolist(), go_aspect=aspect, obo_path=obo_path, train_annotations=_train_annotations(f))
        ap, eligible_a = PreparedCafaEvaluator.prepare(y_true=a.labels, y_prob=a.probabilities, go_terms=a.go_terms.tolist(), go_aspect=aspect, obo_path=obo_path, train_annotations=_train_annotations(a))
        if not np.array_equal(eligible_f, eligible_a) or not eligible_f.all():
            raise ValueError("Cache eligibility is inconsistent with prepared CAFA labels")
        fm, fr = _metric_result(fp, plan, batch_size); am, ar = _metric_result(ap, plan, batch_size)
        for metric in METRICS:
            diffs = fr[metric] - ar[metric]; lo, hi = percentile_interval(diffs)
            records["performance"].append({"aspect": aspect, "category": category, "n": n, "metric": metric, "full": fm[metric], "ablation": am[metric], "delta": fm[metric]-am[metric], "ci_lower": lo, "ci_upper": hi, "status": "ok"})
            records["replicates"][f"{aspect}/{category}/{metric}/full"] = fr[metric]
            records["replicates"][f"{aspect}/{category}/{metric}/ablation"] = ar[metric]
    return records


def _sheet(book: Workbook, name: str, headers: list[str], rows: list[list[Any]]) -> None:
    ws = book.create_sheet(name)
    ws.append(headers)
    for cell in ws[1]: cell.font = Font(bold=True)
    for row in rows: ws.append(row)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = min(48, max(12, max(len(str(c.value or "")) for c in column) + 2))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float): cell.number_format = "0.000000"


def write_workbook(records: dict[str, list[dict[str, Any]]], path: Path, metadata: dict[str, Any]) -> Path:
    book = Workbook(); book.remove(book.active)
    _sheet(book, "README", ["field", "value"], [[k, json.dumps(v) if isinstance(v, (dict, list)) else v] for k, v in metadata.items()])
    counts = records["counts"]
    _sheet(book, "Stratum_counts", ["GO aspect", "category", "display label", "raw N", "evaluated N", "excluded N", "exclusion reason"], [[r["aspect"], r["category"], r["display_label"], r["raw_n"], r["evaluated_n"], r["excluded_n"], r["exclusion_reason"]] for r in counts])
    gates = records["gate"]
    _sheet(book, "Gate_summary", ["GO aspect", "category", "order", "display label", "N evaluated", "model ID", "mean homology weight", "CI lower", "CI upper", "lower error", "upper error", "B requested", "B valid", "status"], [[r[k] for k in ("aspect","category","order","display_label","n","model","point","ci_lower","ci_upper","lower_error","upper_error","n_requested","n_valid","status")] for r in gates])
    pairs = records["gate_paired"]
    _sheet(book, "Gate_paired", ["GO aspect", "stratum", "N", "full mean", "ablation mean", "delta (full - ablation)", "CI lower", "CI upper", "status"], [[r[k] for k in ("aspect","category","n","full","ablation","delta","ci_lower","ci_upper","status")] for r in pairs])
    for aspect in ("BP", "MF", "CC"):
        selection = [r for r in gates if r["aspect"] == aspect]
        wide = []
        for category in CATEGORIES:
            sub = {r["model"]: r for r in selection if r["category"] == category}
            f = sub.get("sequence_homology_confidence_gate", {}); a = sub.get("sequence_homology_internal_gate", {})
            wide.append([LABELS[category], f.get("point"), f.get("ci_lower"), f.get("ci_upper"), f.get("lower_error"), f.get("upper_error"), a.get("point"), a.get("ci_lower"), a.get("ci_upper"), a.get("lower_error"), a.get("upper_error")])
        _sheet(book, f"Gate_plot_{aspect}", ["maximum sequence identity to retained hits (%)", "full mean", "full CI low", "full CI high", "full lower error", "full upper error", "ablation mean", "ablation CI low", "ablation CI high", "ablation lower error", "ablation upper error"], wide)
    perf = records["performance"]
    headers = ["GO aspect", "identity stratum", "N evaluated", "metric", "full model", "no-explicit-signals ablation", "delta (full - ablation)", "paired CI lower", "paired CI upper", "status"]
    _sheet(book, "Performance", headers, [[r[k] for k in ("aspect","category","n","metric","full","ablation","delta","ci_lower","ci_upper","status")] for r in perf])
    for aspect in ("BP", "MF", "CC"):
        _sheet(book, f"Performance_{aspect}", headers[1:9], [[r[k] for k in ("category","n","metric","full","ablation","delta","ci_lower","ci_upper")] for r in perf if r["aspect"] == aspect])
    path.parent.mkdir(parents=True, exist_ok=True); book.save(path)
    load_workbook(path, read_only=True).close()
    return path
